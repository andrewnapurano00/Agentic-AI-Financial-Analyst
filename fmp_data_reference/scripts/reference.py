"""Standalone bounded FMP collector; all writes remain below this script's reference root."""
import argparse, csv, hashlib, io, json, os, re, ssl, time, tempfile
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import urlsplit, parse_qsl, urlencode, quote
import requests
from dotenv import dotenv_values
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT.parent / 'FMP_API_References_and_Example_Schemas.xlsx'
DOC = 'https://site.financialmodelingprep.com/developer/docs'
BASE = 'https://financialmodelingprep.com/stable/'
SECTORS = {'AAPL':'Technology','JPM':'Financial Services','JNJ':'Healthcare','AMZN':'Consumer Cyclical','PG':'Consumer Defensive','CAT':'Industrials','XOM':'Energy','NEE':'Utilities','LIN':'Basic Materials','PLD':'Real Estate','GOOGL':'Communication Services'}
# Explicit semantic replacements; only paths actually present in the official catalogue are used.
MAP = {
1:'search-symbol',2:'search-symbol',3:'search-name',5:'search-cik',6:'search-cusip',7:'search-isin',8:'stock-list',9:'etf-list',10:'financial-statement-symbol-list',11:'actively-trading-list',12:'commitment-of-traders-list',13:'cik-list',15:'symbol-change',16:'batch-exchange-quote',17:'index-list',18:'profile',19:'governance-executive-compensation',20:'executive-compensation-benchmark',21:'company-notes',22:'historical-employee-count',23:'employee-count',24:'company-screener',25:'grades',26:'key-executives',28:'market-capitalization',29:'historical-market-capitalization',30:'available-countries',31:'analyst-estimates',35:'stock-peers',36:'exchange-market-hours',37:'all-exchange-market-hours',38:'delisted-companies',39:'shares-float',41:'shares-float-all',42:'available-sectors',43:'available-industries',44:'available-exchanges',45:'quote',47:'quote-short',49:'batch-exchange-quote',50:'stock-price-change',51:'aftermarket-trade',52:'aftermarket-quote',53:'batch-aftermarket-quote',54:'batch-aftermarket-trade',
61:'income-statement',62:'income-statement',63:'balance-sheet-statement',64:'balance-sheet-statement',65:'cash-flow-statement',66:'cash-flow-statement',67:'income-statement-as-reported',68:'balance-sheet-statement-as-reported',69:'cash-flow-statement-as-reported',70:'financial-statement-full-as-reported',71:'financial-reports-dates',72:'financial-reports-json',73:'financial-reports-xlsx',74:'key-metrics',75:'key-metrics-ttm',76:'ratios',77:'ratios-ttm',78:'cash-flow-statement-growth',79:'income-statement-growth',80:'balance-sheet-statement-growth',81:'financial-growth',82:'financial-scores',83:'owner-earnings',84:'enterprise-values',85:'discounted-cash-flow',86:'custom-discounted-cash-flow',87:'custom-levered-discounted-cash-flow',88:'ratings-snapshot',89:'ratings-historical',91:'price-target-summary',94:'price-target-consensus',96:'grades',97:'grades',98:'grades-consensus',100:'fmp-articles',101:'news/general-latest',102:'news/stock',104:'news/forex',105:'news/crypto',106:'news/press-releases-latest',107:'news/press-releases',111:'earning-call-transcript',112:'earning-call-transcript-dates',
114:'sec-filings-financials',115:'sec-filings-financials',116:'sec-filings-financials',117:'sec-filings-8k',118:'sec-filings-search/symbol',119:'industry-classification-search',120:'all-industry-classification',121:'standard-industrial-classification-list',122:'earnings-calendar',123:'earnings',125:'earnings',126:'dividends-calendar',127:'dividends',128:'splits-calendar',129:'splits',130:'ipos-disclosure',131:'ipos-prospectus',132:'ipos-calendar',133:'mergers-acquisitions-latest',134:'mergers-acquisitions-search',135:'historical-chart/5min',136:'historical-price-eod/full',137:'technical-indicators/sma',138:'technical-indicators/ema',139:'technical-indicators/wma',140:'technical-indicators/dema',141:'technical-indicators/tema',142:'technical-indicators/williams',143:'technical-indicators/rsi',144:'technical-indicators/adx',145:'technical-indicators/standarddeviation',
146:'funds/disclosure-dates',147:'funds/disclosure',148:'etf/holdings',149:'etf/info',150:'etf/sector-weightings',151:'etf/country-weightings',152:'etf/asset-exposure',153:'funds/disclosure-dates',154:'funds/disclosure',155:'funds/disclosure-holders-search',156:'funds/disclosure-holders-latest',157:'esg-disclosures',158:'esg-ratings',159:'esg-benchmark',160:'senate-trades',161:'senate-latest',162:'house-trades',163:'house-latest',164:'batch-index-quotes',165:'sector-pe-snapshot',166:'industry-pe-snapshot',167:'sector-performance-snapshot',168:'historical-sector-performance',169:'biggest-gainers',170:'biggest-losers',171:'most-actives',172:'commitment-of-traders-analysis',173:'commitment-of-traders-analysis',174:'commitment-of-traders-report',175:'commitment-of-traders-report',176:'institutional-ownership/extract',177:'institutional-ownership/dates',179:'institutional-ownership/dates',182:'institutional-ownership/dates',183:'institutional-ownership/latest',184:'institutional-ownership/symbol-positions-summary',186:'institutional-ownership/holder-performance-summary',187:'institutional-ownership/industry-summary',189:'institutional-ownership/extract',191:'insider-trading/latest',192:'insider-trading/search',193:'insider-trading/search',194:'insider-trading-transaction-type',196:'insider-trading/statistics',200:'acquisition-of-beneficial-ownership',202:'crowdfunding-offerings-latest',203:'crowdfunding-offerings-search',204:'crowdfunding-offerings',205:'fundraising-latest',206:'fundraising-search',207:'fundraising',208:'treasury-rates',209:'economic-indicators',210:'economic-calendar',211:'market-risk-premium',212:'commodities-list',213:'batch-commodity-quotes',214:'quote',215:'historical-chart/5min',216:'historical-price-eod/full',217:'forex-list',218:'batch-forex-quotes',219:'quote',220:'historical-chart/5min',221:'historical-price-eod/full',222:'cryptocurrency-list',223:'batch-crypto-quotes',224:'quote',225:'historical-chart/5min',226:'historical-price-eod/full',227:'sp500-constituent',228:'historical-sp500-constituent',229:'nasdaq-constituent',230:'historical-nasdaq-constituent',231:'dowjones-constituent',232:'historical-dowjones-constituent',233:'revenue-product-segmentation',234:'revenue-geographic-segmentation',238:'batch-quote',239:'eod-bulk',240:'income-statement-bulk',241:'balance-sheet-statement-bulk',242:'cash-flow-statement-bulk',245:'earnings-surprises-bulk',246:'profile-bulk',247:'peers-bulk',248:'rating-bulk',249:'dcf-bulk',250:'key-metrics-ttm-bulk',251:'ratios-ttm-bulk',252:'scores-bulk',254:'income-statement-growth-bulk',255:'balance-sheet-statement-growth-bulk',256:'cash-flow-statement-growth-bulk',257:'price-target-summary-bulk',258:'upgrades-downgrades-consensus-bulk',259:'etf-holder-bulk'}
STOCK_IDS = set([2,3,5,6,7,18,19,21,22,23,25,26,27,28,29,31,32,33,34,35,39,40,45,46,47,50,51,52,53,54,57]+list(range(61,100))+[102,103,107,108,111,112,113,118,119,123,125,127,129]+list(range(135,146))+[152,156,157,158,160,162,184,185,188,190,192,193,195,196,199,200,201,233,234])
BATCH_IDS={53,54,238}
STOCK_IDS.difference_update({92,93,95,99})
STOCK_IDS.difference_update({103,193})
STOCK_IDS.update({1,4})
# Latest fund-holder listing is not a verified per-stock substitute.
MAP.pop(156,None)
STATEMENTS=set(range(61,70))
SEMANTIC_NOTES={1:'Search-symbol is a narrower replacement, not a guarantee of identical general instrument/company search coverage.',16:'Batch exchange quotes replace symbol enumeration; not equivalent record coverage.',36:'Exchange market hours; holiday dates are separately documented as holidays-by-exchange.',62:'Original CIK/CUSIP route replaced with a symbol-resolved request; no assertion of native CIK/CUSIP statement support.',64:'Original CIK/CUSIP route replaced with a symbol-resolved request.',66:'Original CIK/CUSIP route replaced with a symbol-resolved request.',97:'Symbol-filtered grades replacement, not an equivalent global RSS feed.',114:'Latest financial filings replacement; not an identical historical RSS stream.',115:'Latest financial filings replacement; not an identical RSS stream.',116:'Financial-filings replacement; not an assertion of identical all-form coverage.',125:'Earnings records expose surprise inputs where present; no fabricated calculated surprise.',146:'Funds disclosure dates; ETF support is tested rather than assumed.',147:'Funds disclosure replacement; holdings date/filing semantics must be checked.',148:'Current ETF holdings endpoint; not an identical legacy report.',156:'Latest fund holders is a market-wide alternative, not equivalent legacy per-stock coverage.'}

def now(): return datetime.now(timezone.utc).isoformat()
def rel(p): return str(p.relative_to(ROOT)).replace('\\','/')
def safe_path(p):
    p=Path(p).resolve()
    if not p.is_relative_to(ROOT): raise ValueError('Output outside reference folder')
    p.parent.mkdir(parents=True,exist_ok=True)
    return p
def dump(p,obj):
    p=safe_path(p)
    with tempfile.NamedTemporaryFile(mode='w',encoding='utf-8',dir=p.parent,suffix='.tmp',delete=False) as f:
        f.write(json.dumps(obj,ensure_ascii=False,indent=2));f.flush();os.fsync(f.fileno());temporary=f.name
    os.replace(temporary,p)
def writecsv(p,rows,fields=None):
    rows=list(rows); fields=fields or list(dict.fromkeys(k for r in rows for k in r))
    p=safe_path(p)
    with tempfile.NamedTemporaryFile(mode='w',encoding='utf-8',newline='',dir=p.parent,suffix='.tmp',delete=False) as f:
        w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)
        f.flush();os.fsync(f.fileno());temporary=f.name
    os.replace(temporary,p)
def key():
    value=os.environ.get('FMP_API_KEY') or dotenv_values(ROOT.parent/'.env').get('FMP_API_KEY')
    if not value:
        p=ROOT.parent/'.streamlit/secrets.toml'
        if p.exists():
            import tomllib
            value=tomllib.loads(p.read_text(encoding='utf-8')).get('FMP_API_KEY')
    if not value: raise SystemExit('Set FMP_API_KEY in your execution environment. Do not paste it into chat.')
    return value
def sanitize(s,secret):
    s=str(s).replace(secret,'[REDACTED]')
    return re.sub(r'(?i)(apikey|api_key|token|access_token|authorization)([=:%\\\s]+)[^&\s"<>]+',r'\1\2[REDACTED]',s)
def workbook():
    if not SOURCE.exists() and (ROOT/'documentation/workbook_tabs.json').exists():
        return json.loads((ROOT/'documentation/workbook_tabs.json').read_text(encoding='utf-8'))
    w=load_workbook(SOURCE,read_only=True,data_only=True)
    tabs={}
    for sheet in w:
        rows=list(sheet.iter_rows(values_only=True)); h=next(i for i,r in enumerate(rows) if r and r[0]=='API ID')
        headers=[str(x) for x in rows[h]]
        tabs[sheet.title]=[dict(zip(headers,r)) for r in rows[h+1:] if r and r[0]]
    w.close()
    dump(ROOT/'documentation/workbook_tabs.json',tabs)
    return tabs
def inventory(tabs):
    cat=json.loads((ROOT/'documentation/official_endpoint_catalogue.json').read_text(encoding='utf-8'))
    official=defaultdict(list)
    for r in cat: official[urlsplit(r['url']).path.removeprefix('/stable/')].append(r)
    out=[]
    for r in tabs['API Reference']:
        n=int(r['API ID'].split('-')[1]); path=MAP.get(n)
        if path and path not in official: raise ValueError('Replacement missing official documentation: '+path)
        docexample=official[path][0]['url'] if path else ''
        out.append({'api_id':r['API ID'],'category':r['Category'],'endpoint_name':r['API Name'],'original_endpoint':r['Endpoint Template'],'original_example':r['Endpoint Example'],'original_version':r['Version'],'workbook_description':r['Description'],'workbook_status':r['Status'],'verified_endpoint':BASE+path if path else '', 'official_example':docexample,'official_documentation':DOC if path else '', 'mapping_status':'official_catalogue_verified_replacement' if path else 'no_current_replacement_verified','mapping_notes':SEMANTIC_NOTES.get(n,'Replacement selected by matching endpoint purpose; current field compatibility must be checked.' if path else 'No exact stable equivalent verified; original workbook route will be tested where HTTP sampling is possible.'),'parameters_verification':'Parameter names/values in official_example are documented. Additional sampling controls originate from workbook or live probes; not an exhaustive required/optional contract.','stock_applicable':n in STOCK_IDS or n in BATCH_IDS,'collection_scope':'stock' if n in STOCK_IDS or n in BATCH_IDS else 'alternative_or_marketwide'})
    writecsv(ROOT/'endpoint_inventory.csv',out)
    logo=next(r for r in out if r['api_id']=='FMP-033')
    logo.update(verified_endpoint='https://financialmodelingprep.com/image-stock/{symbol}.png',official_example='https://financialmodelingprep.com/image-stock/EURUSD.png',official_documentation='https://site.financialmodelingprep.com/developer/docs/company-image-api',mapping_status='official_legacy_documentation_correction',mapping_notes='Official legacy logo example includes .png; workbook omitted extension. Preserve original and test corrected image route.')
    writecsv(ROOT/'endpoint_inventory.csv',out)
    return out,official

class Collector:
    def __init__(self,args):
        self.args=args;self.secret=key();self.session=requests.Session();self.session.headers['apikey']=self.secret
        self.session.max_redirects=0
        self.run=datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ');self.last=0;self.calls=0
        self.cache=ROOT/'raw/request_cache.json'
        try:self.index=json.loads(self.cache.read_text(encoding='utf-8')) if self.cache.exists() else {}
        except (ValueError,UnicodeDecodeError):
            self.index={}
            for metadata in (ROOT/'raw').rglob('*.metadata.json'):
                try:record=json.loads(metadata.read_text(encoding='utf-8'))
                except (ValueError,UnicodeDecodeError):continue
                if record.get('request_signature'):self.index[record['request_signature']]=record
        self.profiles={};self.coverage=[]
        self.endpoint_times={}
        for record in self.index.values():
            try:stamp=datetime.fromisoformat(record['retrieved_at_utc']).timestamp()
            except (KeyError,ValueError):continue
            self.endpoint_times[record['endpoint']]=max(stamp,self.endpoint_times.get(record['endpoint'],0))
        self.bundle=ROOT/'documentation/windows_ca_bundle.pem'
        if not self.bundle.exists(): safe_path(self.bundle).write_text(''.join(ssl.DER_cert_to_PEM_cert(c) for c in ssl.create_default_context().get_ca_certs(binary_form=True)),encoding='ascii')
    def request(self,endpoint,params,api_id):
        params={k:v for k,v in params.items() if k.lower() not in ['apikey','api_key'] and v is not None}
        signature=hashlib.sha256(json.dumps([endpoint,params,self.args.max_bytes],sort_keys=True).encode()).hexdigest()
        if signature in self.index and not self.args.refresh:
            cached=self.index[signature]; body=(ROOT/cached['raw_file']).read_bytes() if cached.get('raw_file') else b''
            if self.secret.encode() in body: raise ValueError('Cached credential reflection blocked')
            return {**cached,'cache_reused':True},body
        result={ 'endpoint':endpoint,'parameters':params,'retrieved_at_utc':now(),'http_status':'','status':'failed','raw_file':'','response_bytes':0,'bytes_limit':self.args.max_bytes,'sampling_limit':8 if str(params.get('limit'))=='8' else 5,'cache_reused':False,'request_signature':signature,'attempts':0,'detail':'' }
        if urlsplit(endpoint).hostname!='financialmodelingprep.com' or not endpoint.startswith('https://'):
            result.update(status='unsupported_transport',detail='Only HTTPS FMP host sampling is supported; no WebSocket subscription opened.')
            return result,b''
        body=b''
        for attempt in range(self.args.retries+1):
            heavy=any(s in endpoint for s in ('-bulk','_bulk','batch-historical-eod'))
            spacing=max(self.args.interval,60 if any(s in endpoint for s in ('profile-bulk','etf-holder-bulk')) else 10 if heavy else self.args.interval)
            time.sleep(max(0,spacing-(time.monotonic()-self.last),spacing-(time.time()-self.endpoint_times.get(endpoint,0))));self.last=time.monotonic();self.calls+=1
            self.endpoint_times[endpoint]=time.time()
            result['attempts']=attempt+1
            try:
                with self.session.get(endpoint,params=params,timeout=(10,self.args.timeout),verify=str(self.bundle),stream=True,allow_redirects=False) as response:
                    result['http_status']=response.status_code;result['content_type']=response.headers.get('Content-Type','')
                    chunks=[];size=0;truncated=False
                    for chunk in response.iter_content(65536):
                        if size+len(chunk)>self.args.max_bytes:
                            remaining=self.args.max_bytes-size;chunks.append(chunk[:remaining]);size+=remaining;truncated=True;break
                        chunks.append(chunk);size+=len(chunk)
                    body=b''.join(chunks)
                    if self.secret.encode() in body:
                        result.update(status='credential_reflection_quarantined',detail='Response contained configured credential and was not saved.');body=b'';break
                    result.update(response_bytes=len(body),retrieved_at_utc=now(),body_truncated=truncated)
                    result['status']='bounded_download' if truncated else 'success' if response.ok else 'restricted' if response.status_code in (401,402,403) else 'not_found' if response.status_code==404 else 'rate_limited' if response.status_code==429 else 'failed'
                    if not truncated:
                        try:
                            payload=json.loads(body)
                            error_object=payload if isinstance(payload,dict) else payload[0] if isinstance(payload,list) and len(payload)==1 and isinstance(payload[0],dict) else {}
                            err=error_object.get('Error Message') or error_object.get('error') or error_object.get('Error')
                            if err:
                                result['detail']=sanitize(err,self.secret)[:600]
                                result['status']='deprecated' if 'legacy' in str(err).lower() else 'restricted' if any(x in str(err).lower() for x in ['premium','subscription','upgrade','restricted','plan','invalid api']) else 'api_error'
                            elif response.ok and payload in ([],{},None):result['status']='empty'
                        except (ValueError,UnicodeDecodeError):
                            text=body.decode('utf-8',errors='replace')
                            if any(x in text.lower() for x in ['error message','upgrade your','invalid api key','premium query','legacy endpoint']):
                                result['status']='deprecated' if 'legacy endpoint' in text.lower() else 'restricted';result['detail']=sanitize(text,self.secret)[:600]
                    if response.status_code in (429,500,502,503,504) and attempt<self.args.retries:
                        time.sleep(min(30,2**(attempt+1)));continue
                    break
            except requests.RequestException as e:
                result.update(status='transport_failed',detail=type(e).__name__)
                if attempt<self.args.retries:time.sleep(2**(attempt+1));continue
                break
        if body:
            extension='.partial' if result.get('body_truncated') else '.body'
            p=ROOT/'raw'/api_id/self.run/(signature[:16]+extension)
            safe_path(p).write_bytes(body);result['raw_file']=rel(p)
            result['sha256']=hashlib.sha256(body).hexdigest()
        dump(ROOT/'raw'/api_id/self.run/(signature[:16]+'.metadata.json'),result)
        self.index[signature]=result;dump(self.cache,self.index)
        return result,body
    def profiles_first(self):
        rows=[]
        for ticker,expected in SECTORS.items():
            meta,body=self.request(BASE+'profile',{'symbol':ticker},'FMP-018')
            try:data=json.loads(body)
            except (ValueError,UnicodeDecodeError):data=[]
            profile=next((r for r in data if isinstance(r,dict) and r.get('symbol')==ticker),{}) if isinstance(data,list) else {}
            self.profiles[ticker]=profile
            rows.append({'ticker':ticker,'requested_sector':expected,'returned_sector':profile.get('sector',''),'returned_industry':profile.get('industry',''),'returned_currency':profile.get('currency',''),'cik':profile.get('cik',''),'cusip':profile.get('cusip',''),'isin':profile.get('isin',''),'retrieved_at_utc':meta['retrieved_at_utc'],'status':meta['status'] if profile else 'profile_unavailable','classification_notes':'Exact label match' if profile.get('sector')==expected else 'Returned FMP label differs; original selection retained, not silently relabeled.' if profile else 'Sector could not be verified; selection retained with disclosure.','raw_file':meta['raw_file']})
        writecsv(ROOT/'sector_tickers.csv',rows)
    def params(self,item,ticker,variant,official):
        n=int(item['api_id'][4:]);path=None if item.get('_legacy_sample') else MAP.get(n); original=urlsplit(item['original_example']);p=dict(parse_qsl(urlsplit(item['official_example']).query)) if path else dict(parse_qsl(original.query))
        endpoint=item['verified_endpoint'] or item['original_endpoint']
        for param in list(p):
            if param.lower() in ['apikey','api_key']:p.pop(param)
        p.pop('page',None)
        if ticker in SECTORS:
            if path:
                if 'symbols' in p or n in BATCH_IDS:p['symbols']=','.join(SECTORS) if n in BATCH_IDS else ticker;p.pop('symbol',None)
                elif path=='search-symbol':p['query']=ticker
                elif path=='search-name':p['query']=self.profiles[ticker].get('companyName',ticker)
                elif path=='search-cik':p['cik']=self.profiles[ticker].get('cik')
                elif path=='search-cusip':p['cusip']=self.profiles[ticker].get('cusip')
                elif path=='search-isin':p['isin']=self.profiles[ticker].get('isin')
                elif path=='industry-classification-search':p['symbol']=ticker
                elif path=='funds/disclosure-holders-latest':pass
                else:p['symbol']=ticker
            else:
                for placeholder in ['symbol','symbols']:
                    endpoint=endpoint.replace('{'+placeholder+'}',ticker)
                endpoint=endpoint.replace('{company_name}',quote(self.profiles[ticker].get('companyName',ticker),safe=''))
                for param in ['symbol','symbols','ticker','tickers']:
                    if param in p:p[param]=ticker
                if '{cik_or_cusip}' in endpoint:endpoint=endpoint.replace('{cik_or_cusip}',str(self.profiles[ticker].get('cik','')))
        else:
            for placeholder in re.findall(r'\{([^}]+)\}',endpoint):
                parts=original.path.split('/');template_parts=urlsplit(endpoint).path.split('/');i=next((j for j,v in enumerate(template_parts) if '{'+placeholder+'}' in v),-1)
                example=parts[i] if 0<=i<len(parts) else ''
                endpoint=endpoint.replace('{'+placeholder+'}',example)
            if n in (214,215,216):p['symbol']='GCUSD'
            if n in (219,220,221):p['symbol']='EURUSD'
            if n in (224,225,226):p['symbol']='BTCUSD'
            if path and path.startswith('historical') and 'symbol' in p and n in (214,215,216,219,220,221,224,225,226):pass
        # Bound only documented/workbook parameter names, plus empirically checked statement limit/period controls.
        if 'limit' in p:p['limit']=str(8 if variant=='quarter' else 5)
        if n in STATEMENTS:p.update(period=variant,limit=str(8 if variant=='quarter' else 5))
        if path and n in (74,76,78,79,80,81,84,233,234):p.update(period='annual',limit='5')
        if path and (path.startswith('historical-chart') or path.startswith('historical-price-eod')):
            end=datetime.now(timezone.utc).date();p.update({'from':str(end-timedelta(days=12 if 'price-eod' in path else 2)),'to':str(end)})
        if path and path.startswith('technical-indicators/'):
            p.setdefault('periodLength','10');p.setdefault('timeframe','5min');end=datetime.now(timezone.utc).date();p.update({'from':str(end-timedelta(days=21)),'to':str(end)})
        if 'from' in p and not (path and ('historical-chart' in path or 'historical-price-eod' in path or 'technical-indicators' in path)):
            end=datetime.now(timezone.utc).date();p.update({'from':str(end-timedelta(days=14)),'to':str(end)})
        if path and path.endswith('-snapshot'):p['date']=str(datetime.now(timezone.utc).date()-timedelta(days=1))
        if path=='eod-bulk':p['date']=str(datetime.now(timezone.utc).date()-timedelta(days=1))
        if path and 'year' in p:p['year']=str(datetime.now(timezone.utc).year-1)
        if path and 'period' in p and '-bulk' in path:p['period']='Q4'
        return endpoint,p
    def collect_item(self,item,official):
        n=int(item['api_id'][4:]);targets=list(SECTORS) if item['stock_applicable'] else ['alternative']
        variants=['annual','quarter'] if n in STATEMENTS else ['annual'] if n in (74,76,78,79,80,81,84,233,234) else ['default']
        for ticker in targets:
            for variant in variants:
                endpoint,params=self.params(item,ticker,variant,official)
                meta,body=self.request(endpoint,params,item['api_id'])
                sample=None;status=meta['status'];kind='';rows=0;dates=[]
                if status in ['success','empty','bounded_download']:
                    try:
                        payload=json.loads(body);kind='json'
                    except (ValueError,UnicodeDecodeError):
                        text=body.decode('utf-8',errors='replace')
                        content_type=meta.get('content_type','').lower()
                        if body.startswith(b'PK') or 'spreadsheet' in content_type or 'excel' in content_type or 'image/' in content_type:
                            payload=None;kind='binary_response'
                        elif 'csv' in content_type or (',' in text.partition('\n')[0] and not text.lstrip().startswith(('<','[','{'))):
                            if status=='bounded_download':text=text.rpartition('\n')[0]
                            csv.field_size_limit(self.args.max_bytes)
                            payload=list(csv.DictReader(io.StringIO(text)));kind='csv_strings'
                        elif status=='bounded_download' and text.lstrip().startswith('['):
                            payload=[];decoder=json.JSONDecoder();cursor=text.find('[')+1
                            while cursor<len(text):
                                while cursor<len(text) and text[cursor] in ' \r\n\t,':cursor+=1
                                try:v,end=decoder.raw_decode(text,cursor)
                                except ValueError:break
                                payload.append(v);cursor=end
                            kind='json_complete_records_from_bounded_prefix'
                        else:payload=None;kind='binary_or_text'
                    if payload is not None:
                        if isinstance(payload,list) and (n in BATCH_IDS or item['category']=='Bulk Requests') and ticker!='alternative':
                            payload=[r for r in payload if isinstance(r,dict) and r.get('symbol')==ticker]
                        if item['category']=='Bulk Requests' and isinstance(payload,list):
                            matched=[r for r in payload if isinstance(r,dict) and r.get('symbol') in SECTORS]
                            if matched:payload=matched
                        limit=8 if variant=='quarter' else 5
                        sample=bounded_sample(payload,limit)
                        rows=len(sample) if isinstance(sample,list) else 1 if sample else 0
                        if rows==0:status='empty'
                        dates=extract_dates(sample)
                samplefile='';binaryfile=''
                if sample is not None:
                    dest=ROOT/'samples'/item['api_id']/self.run/(ticker+'-'+('legacy-' if item.get('_legacy_sample') else '')+variant+'.json');dump(dest,sample);samplefile=rel(dest)
                elif status in ('success','bounded_download') and kind in ('binary_or_text','binary_response'):
                    extension='.png' if 'image/png' in meta.get('content_type','') else '.xlsx' if body.startswith(b'PK') else '.bin'
                    if status=='bounded_download':extension='.partial'+extension
                    dest=ROOT/'samples'/item['api_id']/self.run/(ticker+'-'+variant+extension);safe_path(dest).write_bytes(body);binaryfile=rel(dest)
                    if status=='success':status='success_non_json'
                self.coverage.append({'api_id':item['api_id'],'category':item['category'],'endpoint_name':item['endpoint_name'],'ticker_or_input':ticker,'variant':('legacy-' if item.get('_legacy_sample') else '')+variant,'status':status,'http_status':meta['http_status'],'endpoint':endpoint,'parameters_json':json.dumps(params,sort_keys=True),'retrieved_at_utc':meta['retrieved_at_utc'],'record_count':rows,'data_dates_json':json.dumps(dates),'sample_limit':8 if variant=='quarter' else 5,'raw_file':meta['raw_file'],'sample_file':samplefile,'format':kind,'cache_reused':meta['cache_reused'],'detail':meta['detail'],'applicability_notes':'Sector-stock request' if ticker in SECTORS else 'Non-stock/universe/bulk input; sector symbols are not independent requests. Documented/example parameters listed explicitly.','parameter_controls_basis':'Official example + workbook; additional period/limit/date sampling controls live-probed, not fully documented required rules.'})
                self.coverage[-1]['binary_sample_file']=binaryfile
                if item['category']=='Bulk Requests' and ticker=='alternative' and sample is not None and isinstance(sample,list):
                    for stock in SECTORS:
                        records=[r for r in payload if isinstance(r,dict) and r.get('symbol')==stock]
                        subset=bounded_sample(records,5);dest=ROOT/'samples'/item['api_id']/self.run/(stock+'-bulk.json');dump(dest,subset)
                        self.coverage.append({**self.coverage[-1],'ticker_or_input':stock,'sample_file':rel(dest),'record_count':len(subset),'status':'success' if subset else 'not_present_in_bounded_bulk','applicability_notes':'Extracted locally from single bounded bulk request; absent symbols may be outside downloaded partition.'})
        unique={}
        for row in self.coverage:
            unique[(row['api_id'],row['ticker_or_input'],row['variant'],row['endpoint'],row['parameters_json'])]=row
        self.coverage=list(unique.values())
        writecsv(ROOT/'coverage_matrix.csv',self.coverage)
        writecsv(ROOT/'raw'/('coverage-'+self.run+'.csv'),self.coverage)
        dump(ROOT/'raw'/('run-'+self.run+'.json'),{'run':self.run,'updated_at_utc':now(),'network_requests':self.calls,'complete_api_count':len(set(r['api_id'] for r in self.coverage)),'settings':{'interval_seconds':self.args.interval,'timeout_seconds':self.args.timeout,'bounded_retries':self.args.retries,'max_bytes':self.args.max_bytes,'pagination_pages':1,'refresh':self.args.refresh}})

def bounded_sample(value,limit):
    if isinstance(value,list):return [bounded_sample(v,limit) for v in value[:limit]]
    if isinstance(value,dict):return {k:bounded_sample(v,limit) for k,v in value.items()}
    return value
def extract_dates(value):
    found=set()
    def visit(v):
        if isinstance(v,dict):
            for k,x in v.items():
                if ('date' in k.lower() or 'time' in k.lower() or k.lower() in ['period','fiscalyear','calendaryear']) and isinstance(x,(str,int)):found.add(str(x))
                if isinstance(x,(dict,list)):visit(x)
        elif isinstance(v,list):
            for x in v:visit(x)
    visit(value);return sorted(found)[:30]

def main():
    p=argparse.ArgumentParser();p.add_argument('--smoke',action='store_true');p.add_argument('--inventory-only',action='store_true');p.add_argument('--refresh',action='store_true');p.add_argument('--legacy-fallback',action='store_true');p.add_argument('--ids',nargs='*');p.add_argument('--interval',type=float,default=0.75);p.add_argument('--timeout',type=float,default=25);p.add_argument('--retries',type=int,default=1);p.add_argument('--max-bytes',type=int,default=3*1024*1024);a=p.parse_args()
    if a.interval<0.25 or not 1<=a.timeout<=60 or not 0<=a.retries<=3 or not 1024<=a.max_bytes<=16*1024*1024:raise SystemExit('Unsafe collection bounds')
    tabs=workbook();items,official=inventory(tabs)
    if a.inventory_only:print('Inventory',len(items),'tabs',len(tabs));return
    c=Collector(a);c.profiles_first()
    chosen=[x for x in items if (not a.ids or x['api_id'] in a.ids) and (not a.smoke or x['api_id'] in ['FMP-018','FMP-045','FMP-061','FMP-208'])]
    previous=list(csv.DictReader((ROOT/'coverage_matrix.csv').open(encoding='utf-8'))) if (ROOT/'coverage_matrix.csv').exists() else []
    c.coverage=previous
    if a.legacy_fallback:
        blocked_ids={r['api_id'] for r in previous if r['status'] in ('restricted','not_found','api_error','failed') and not r['variant'].startswith('legacy-')}
        chosen=[{**x,'verified_endpoint':'','official_example':'','_legacy_sample':True} for x in chosen if x['api_id'] in blocked_ids and x['verified_endpoint'] and '/api/v' in x['original_endpoint']]
    for i,item in enumerate(chosen,1):
        c.collect_item(item,official)
        counts=Counter(r['status'] for r in c.coverage);print(i,len(chosen),item['api_id'],'requests',c.calls,'statuses',dict(counts),flush=True)
    print('Complete',len(chosen),'references','network_requests',c.calls,flush=True)

if __name__=='__main__':
    try:main()
    except Exception as exc:
        # Do not surface request exception repr or traceback, which may contain sensitive headers/URLs.
        print('Collection stopped:',type(exc).__name__,'Resume with the same script; cached responses are preserved.',flush=True)
        raise SystemExit(1)
