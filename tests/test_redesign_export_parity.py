"""Offline checks of saved download bytes and CSV text safety."""

import csv
import io
import json
from unittest.mock import Mock

import pandas as pd
import pdfplumber
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from langgraphagenticai.ui import ai_portfolio_manager_tab as manager_ui
from langgraphagenticai.ui import stock_screener_tab as screener_ui
from langgraphagenticai.ui import equity_report_tab as equity_ui
from langgraphagenticai.ui.guided_research_tab import guided_result_json
from langgraphagenticai.ui import guided_research_tab as guided_ui
from langgraphagenticai.research.guided_workflow import config as guided_config, fingerprint
from langgraphagenticai.deep_research.presentation import build_research_pdf, clean_report_markdown
from langgraphagenticai.deep_research.models import dumps as deep_dumps
from langgraphagenticai.analysis.investment_brief import exports as brief_exports
from langgraphagenticai.utils.export_csv import safe_csv_bytes


def _csv(data):
    return list(csv.DictReader(io.StringIO(data.decode("utf-8") if isinstance(data, bytes) else data)))


def test_csv_export_text_is_inert_redacted_and_numeric_cells_remain_numeric():
    frame = pd.DataFrame([{
        "Ticker": "AAPL", "Company Name": " \t=1+2", "Note": "\t+SUM(1,2)",
        "Other Text": " @unsafe", "Provider Link": "https://example.test/?token=synthetic-secret",
        "Market Cap": -20.5, "Zero": 0.0, "Mixed": -20.5, "Missing": float("nan"),
        "Nested": {"api_key": "nested-secret", "source": "https://example.test/?token=nested-secret"},
    }])
    original = frame.copy(deep=True)
    data = safe_csv_bytes(frame)
    row = _csv(data)[0]
    assert list(row) == list(frame)
    assert row["Company Name"] == "' \t=1+2"
    assert row["Note"] == "'\t+SUM(1,2)"
    assert row["Other Text"] == "' @unsafe"
    assert "synthetic-secret" not in data.decode("utf-8")
    assert "nested-secret" not in data.decode("utf-8")
    assert row["Market Cap"] == "-20.5" and row["Zero"] == "0.0"
    assert row["Mixed"] == "-20.5" and row["Missing"] == ""
    pd.testing.assert_frame_equal(frame, original)


def test_categorical_export_text_is_inert_in_csv_and_actual_equity_xlsx():
    frame = pd.DataFrame({
        "Ticker": ["AAPL"], "Final Rank": [1],
        "Category Text": pd.Categorical(["=1+2"]),
        "Category Secret": pd.Categorical(["https://example.test/?token=categorical-secret"]),
        "Numeric Fact": [-20.5], "Zero": [0],
    })
    original = frame.copy(deep=True)
    row = _csv(safe_csv_bytes(frame))[0]
    assert row["Category Text"] == "'=1+2"
    assert "categorical-secret" not in str(row)
    assert row["Numeric Fact"] == "-20.5" and row["Zero"] == "0"
    excel = equity_ui.generate_excel_report(
        frame, pd.DataFrame(), {}, pd.DataFrame(), pd.DataFrame(), pd.DataFrame(),
        executive_summary="Saved analysis", peer_group="General / Cross-Sector",
    )
    sheet = load_workbook(io.BytesIO(excel), data_only=False)["Full Scorecard"]
    assert sheet["C2"].data_type == "s" and sheet["C2"].value == "'=1+2"
    assert "categorical-secret" not in sheet["D2"].value
    assert sheet["E2"].data_type == "n" and sheet["E2"].value == -20.5
    assert sheet["F2"].data_type == "n" and sheet["F2"].value == 0
    pd.testing.assert_frame_equal(frame, original)


def test_dynamic_export_header_is_inert_in_csv_and_actual_equity_xlsx():
    frame = pd.DataFrame({"=1+2": ["@unsafe"], "https://example.test/?api_key=header-secret": [0],
                          "Numeric Fact": [-20.5]})
    original = frame.copy(deep=True)
    row = _csv(safe_csv_bytes(frame))[0]
    assert list(row)[0] == "'=1+2" and row["'=1+2"] == "'@unsafe"
    assert "header-secret" not in str(row)
    assert row["Numeric Fact"] == "-20.5"
    excel = equity_ui.generate_excel_report(
        frame, pd.DataFrame(), {}, pd.DataFrame(), pd.DataFrame(), pd.DataFrame(),
        executive_summary="Saved analysis", peer_group="General / Cross-Sector",
    )
    sheet = load_workbook(io.BytesIO(excel), data_only=False)["Full Scorecard"]
    assert sheet["A1"].data_type == "s" and sheet["A1"].value == "'=1+2"
    assert "header-secret" not in sheet["B1"].value
    assert sheet["C2"].data_type == "n" and sheet["C2"].value == -20.5
    pd.testing.assert_frame_equal(frame, original)


def test_saved_screener_csv_matches_raw_result_not_formatted_display(monkeypatch):
    exported = {}
    original_download = screener_ui.st.download_button
    def capture(label, *args, **kwargs):
        exported[label] = kwargs.get("data", args[0] if args else None)
        return original_download(label, *args, **kwargs)
    monkeypatch.setattr(screener_ui.st, "download_button", capture)
    provider = Mock()
    monkeypatch.setattr(screener_ui, "fmp_company_screener_safe", provider)
    frame = pd.DataFrame([{"Ticker": "AAPL", "Company Name": "=1+2", "Market Cap": -20.5, "Price": 0.0}])
    app = AppTest.from_string(
        "from langgraphagenticai.ui.stock_screener_tab import render_stock_screener_tab\nrender_stock_screener_tab('synthetic')",
        default_timeout=30,
    )
    app.session_state["stock_screener_payload"] = {"generated_at": "2026-10-14T12:00:00Z", "results": frame}
    app.run()
    assert not app.exception
    row = _csv(exported["Download Results CSV"])[0]
    assert list(row) == list(frame)
    assert row["Company Name"] == "'=1+2" and row["Market Cap"] == "-20.5"
    assert row["Price"] == "0.0"
    app.run()
    provider.assert_not_called()


def test_manager_saved_recommendation_rebalance_and_template_bytes(monkeypatch):
    exported = {}
    original_download = manager_ui.st.download_button
    def capture(label, *args, **kwargs):
        exported[label] = kwargs.get("data", args[0] if args else None)
        return original_download(label, *args, **kwargs)
    monkeypatch.setattr(manager_ui.st, "download_button", capture)
    run = Mock()
    monkeypatch.setattr(manager_ui, "run_hybrid_portfolio_workflow", run)
    rec = pd.DataFrame([{"ticker": "AAPL", "company_name": "@unsafe", "current_weight": .2,
                         "target_weight": .3, "trade_value": -20.5, "committee_reason": "=1+2"}])
    rebalance = pd.DataFrame([{"ticker": "AAPL", "rebalance_reason": "+SUM(1,2)",
                               "trade_value": -20.5, "target_weight": .3}])
    app = AppTest.from_string(
        "from langgraphagenticai.ui.ai_portfolio_manager_tab import render_ai_portfolio_manager_tab\nrender_ai_portfolio_manager_tab(openai_api_key='synthetic', model_name='synthetic')",
        default_timeout=30,
    )
    app.session_state["pm_hybrid_bundle_v1"] = {
        "portfolio_summary": {}, "run_diagnostics": {}, "recommendation_table": rec,
        "rebalance_table": rebalance, "sector_allocation_table": pd.DataFrame(),
        "evidence_table": pd.DataFrame(), "monitoring_table": pd.DataFrame(),
    }
    app.run()
    assert not app.exception
    assert list(_csv(exported["Download recommendation table CSV"])[0]) == list(rec)
    assert _csv(exported["Download recommendation table CSV"])[0]["committee_reason"] == "'=1+2"
    assert _csv(exported["Download recommendation table CSV"])[0]["trade_value"] == "-20.5"
    assert _csv(exported["Download rebalance table CSV"])[0]["rebalance_reason"] == "'+SUM(1,2)"
    assert list(_csv(exported["Download rebalance table CSV"])[0]) == list(rebalance)
    assert list(_csv(exported["Download holdings template"])[0]) == ["ticker", "shares", "cost_basis"]
    app.run()
    run.assert_not_called()


def test_equity_actual_pdf_and_xlsx_contain_saved_identity_and_numeric_values():
    scorecard = pd.DataFrame([{"Ticker": "AAPL", "Final Rank": 1, "Numeric Fact": -20.5,
                               "Company Name": "=1+2", "Leading": " \t+SUM(1,2)",
                               "At Text": "@unsafe", "Zero": 0.0}])
    original = scorecard.copy(deep=True)
    pdf = equity_ui.generate_story_pdf_report(
        scorecard, {"Profile": pd.DataFrame([{"Ticker": "AAPL", "Fact": "Saved evidence"}])},
        ["AAPL"], "Saved deterministic thesis & risks", None, "General / Cross-Sector",
    )
    with pdfplumber.open(io.BytesIO(pdf)) as document:
        pdf_text = " ".join(page.extract_text() or "" for page in document.pages)
    assert "AAPL" in pdf_text and "Saved deterministic thesis" in pdf_text
    excel = equity_ui.generate_excel_report(
        scorecard, pd.DataFrame(), {}, pd.DataFrame(), pd.DataFrame(), pd.DataFrame(),
        executive_summary="Saved deterministic thesis", peer_group="General / Cross-Sector",
    )
    workbook = load_workbook(io.BytesIO(excel), read_only=False, data_only=False)
    assert {"Executive Summary", "Full Scorecard", "Research Ranking"} <= set(workbook.sheetnames)
    rows = list(workbook["Full Scorecard"].values)
    assert rows[0] == tuple(scorecard.columns)
    assert rows[1][0] == "AAPL" and rows[1][2] == -20.5 and rows[1][6] == 0.0
    for cell in ("D2", "E2", "F2"):
        assert workbook["Full Scorecard"][cell].data_type == "s"
        assert workbook["Full Scorecard"][cell].value.startswith("'")
    assert workbook["Full Scorecard"]["C2"].data_type == "n"
    assert workbook["Full Scorecard"]["G2"].data_type == "n"
    pd.testing.assert_frame_equal(scorecard, original)


def test_guided_saved_json_retains_identity_and_evidence_but_filters_secrets_and_nonfinite():
    result = {
        "run_id": "saved-run", "config_fingerprint": "config-id",
        "evidence_fingerprint": "evidence-id", "answer": {"interpretation": "Saved answer"},
        "evidence": [{"id": "E1", "provider": "FMP", "value": 0,
                      "source_url": "https://example.test/?token=synthetic-secret",
                      "api_key": "synthetic-secret", "missing": float("nan")}],
        "usage": {"model_attempts": 1, "cost": float("inf")},
        "password": "synthetic-secret", "raw_bytes": b"binary",
    }
    encoded = guided_result_json(result)
    packet = json.loads(encoded)
    assert packet["run_id"] == "saved-run"
    assert packet["config_fingerprint"] == "config-id"
    assert packet["evidence_fingerprint"] == "evidence-id"
    assert packet["evidence"][0]["id"] == "E1"
    assert packet["evidence"][0]["value"] == 0
    assert packet["evidence"][0]["missing"] is None
    assert packet["usage"]["cost"] is None
    assert "synthetic-secret" not in encoded and "binary" not in encoded
    assert len(encoded.encode("utf-8")) < 500_000


def test_active_equity_saved_downloads_keep_raw_display_ranking_pdf_and_xlsx(monkeypatch):
    exported = {}
    original_download = equity_ui.st.download_button
    def capture(label, *args, **kwargs):
        exported[label] = kwargs.get("data", args[0] if args else None)
        return original_download(label, *args, **kwargs)
    monkeypatch.setattr(equity_ui.st, "download_button", capture)
    collect = Mock()
    monkeypatch.setattr(equity_ui, "prepare_financials", collect)
    monkeypatch.setattr(equity_ui, "add_equity_research_decision_layer", lambda frame, peer_group: frame)
    monkeypatch.setattr(equity_ui, "add_story_template_fields", lambda frame: frame)
    monkeypatch.setattr(equity_ui, "_build_research_ranking_table", lambda frame: frame[["Ticker", "Final Rank"]].copy())
    monkeypatch.setattr(equity_ui, "build_story_report_tables", lambda frame, peer_group: {})
    monkeypatch.setattr(equity_ui, "build_executive_summary_paragraph", lambda frame: "Saved deterministic thesis.")
    monkeypatch.setattr(equity_ui, "make_display_df", lambda frame: frame.assign(**{"Numeric Fact": ["displayed"]}) if "Numeric Fact" in frame else frame)
    monkeypatch.setattr(equity_ui, "render_sector_metric_coverage_audit", lambda *args: None)
    monkeypatch.setattr(equity_ui, "build_price_return_long_df", lambda *args: pd.DataFrame())
    monkeypatch.setattr(equity_ui, "build_selected_period_return_table", lambda frame: pd.DataFrame())
    scorecard = pd.DataFrame([{"Ticker": "AAPL", "Final Rank": 1, "Numeric Fact": -20.5,
                               "Company Name": "=1+2"}])
    app = AppTest.from_string(
        "from langgraphagenticai.ui.equity_report_tab import render_equity_report_tab\nrender_equity_report_tab(fmp_api_key='synthetic', openai_api_key='', model_name='synthetic')",
        default_timeout=30,
    )
    app.session_state["equity_report_payload"] = {
        "generated_at": "2026-10-14T12:00:00Z", "tickers": ["AAPL"],
        "combined_scorecard": scorecard, "peer_group": "General / Cross-Sector",
        "gpt_note": None, "price_history": {},
    }
    from langgraphagenticai.equity_committee import build_equity_committee_packet
    packet = build_equity_committee_packet(scorecard, peer_group="General / Cross-Sector", generated_at="2026-10-14T12:00:00Z")
    committee_run = Mock()
    monkeypatch.setattr(equity_ui, "run_equity_committee", committee_run)
    app.session_state["equity_report_payload"]["crewai_equity_decision"] = {
        "input_fingerprint": packet["fingerprint"], "best_buy": "AAPL", "confidence": 65,
        "verdict": "WATCH", "summary": "SAVED_COMMITTEE_END", "debate_summary": "Dissent preserved",
    }
    app.run()
    assert not app.exception
    assert _csv(exported["Download Raw CSV"])[0]["Numeric Fact"] == "-20.5"
    assert _csv(exported["Download Raw CSV"])[0]["Company Name"] == "'=1+2"
    assert _csv(exported["Download Display CSV"])[0]["Numeric Fact"] == "displayed"
    assert list(_csv(exported["Download Ranking CSV"])[0]) == ["Ticker", "Final Rank"]
    with pdfplumber.open(io.BytesIO(exported["Download PDF Report"])) as document:
        text = " ".join(page.extract_text() or "" for page in document.pages)
        assert "Saved deterministic thesis" in text
        assert "SAVED_COMMITTEE_END" in text
    workbook = load_workbook(io.BytesIO(exported["Download Excel Research Pack"]), data_only=False)
    assert workbook["Full Scorecard"]["C2"].value == -20.5
    assert workbook["Full Scorecard"]["D2"].data_type == "s"
    next(button for button in app.get("download_button") if button.label == "Download Raw CSV").click().run()
    assert not app.exception
    app.run()
    collect.assert_not_called()
    committee_run.assert_not_called()
    app.session_state["equity_report_payload"]["crewai_equity_decision"]["input_fingerprint"] = "stale"
    app.run()
    assert not app.exception
    with pdfplumber.open(io.BytesIO(exported["Download PDF Report"])) as document:
        text = " ".join(page.extract_text() or "" for page in document.pages)
        assert "SAVED_COMMITTEE_END" not in text
        assert "No completed CrewAI decision matches" in text
    committee_run.assert_not_called()
    collect.assert_not_called()


def test_oversized_guided_export_warns_without_hiding_saved_result_or_calling_models(monkeypatch):
    prepare = Mock()
    run = Mock()
    monkeypatch.setattr(guided_ui, "prepare", prepare)
    monkeypatch.setattr(guided_ui, "run", run)
    question = "What do the business, latest annual financials and recent news establish?"
    cfg = guided_config("AAPL", question, "synthetic", 1.0, None, None, True)
    saved = {"run_id": "saved-run", "config_fingerprint": fingerprint(cfg),
             "plan_fingerprint": "plan-id", "plan": {"steps": []},
             "usage": {"model_attempts": 1, "reserved_cost": None}}
    result = {"run_id": "saved-run", "config_fingerprint": fingerprint(cfg),
              "plan_fingerprint": "plan-id", "evidence": [], "evidence_fingerprint": fingerprint([]),
              "answer": {"interpretation": "Saved interpretation remains visible.", "facts": [],
                         "missing_inputs": [], "risks": []},
              "warnings": [], "steps": [], "usage": saved["usage"],
              "unused_large_field": "X" * 500_001}
    app = AppTest.from_string(
        "from langgraphagenticai.ui.guided_research_tab import render_guided_research_tab\nrender_guided_research_tab(openai_api_key='synthetic', model_name='synthetic')",
        default_timeout=30,
    )
    app.session_state["guided_plan"] = saved
    app.session_state["guided_result"] = result
    app.session_state["guided_widget_values"] = {"guided_unknown_ack": True}
    app.run()
    assert not app.exception
    assert any("Saved interpretation remains visible" in item.value for item in app.markdown)
    assert any("500 KB limit" in item.value for item in app.warning)
    assert not any(button.label == "Download saved evidence and answer" for button in app.get("download_button"))
    app.run()
    prepare.assert_not_called()
    run.assert_not_called()


def test_deep_and_brief_saved_artifacts_keep_identity_status_and_safe_values():
    result = {
        "id": "saved-deep", "request": {"symbols": ["AAPL"], "horizon": "1-3 years"},
        "status": "needs_review", "report": "## Executive Investment Thesis\nSaved caveat [E001].",
        "comparison": [{"Ticker": "AAPL", "Forward EPS": 10, "Notes": "=1+2"}],
        "api_key": "synthetic-secret", "source_url": "https://example.test/?token=synthetic-secret",
        "missing": float("nan"),
    }
    clean_memo = clean_report_markdown(result["report"])
    assert "Saved caveat" in clean_memo
    pdf = build_research_pdf(result)
    with pdfplumber.open(io.BytesIO(pdf)) as document:
        text = " ".join(page.extract_text() or "" for page in document.pages)
    assert "AAPL" in text and "Saved caveat" in text
    deep_packet = json.loads(deep_dumps(result))
    assert deep_packet["id"] == "saved-deep" and deep_packet["status"] == "needs_review"
    assert deep_packet["missing"] is None
    assert "synthetic-secret" not in json.dumps(deep_packet)
    brief_json, brief_csv = brief_exports({"schema": "axiom-saved-brief-v1", "result_id": "saved-deep",
                                            "scenarios": [{"name": "Bear", "change": -20.5, "note": "@unsafe"}]})
    assert json.loads(brief_json)["result_id"] == "saved-deep"
    assert any(row[1] == "'@unsafe" for row in list(csv.reader(io.StringIO(brief_csv)))[1:])
    assert "-20.5" in brief_csv
